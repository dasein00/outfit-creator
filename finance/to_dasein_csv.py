#!/usr/bin/env python3
"""
Заполненная таблица «Расходы_и_доходы_DASEIN.xlsx» → CSV для DASEIN (Финансы → Тетрадь финансов).

    python3 to_dasein_csv.py заполненная.xlsx [выход.csv]

Берёт лист «Записи» (колонки Дата · Расход · Доход · Примечание), пропускает пустые строки,
подставляет дату из строки выше, если она не указана, и печатает, что не удалось разобрать.
"""
import csv
import datetime as dt
import re
import sys

import openpyxl


def amount(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return abs(float(v)) or None
    t = str(v).lower().replace(" ", "").replace(" ", "").replace("₽", "").replace("руб", "").replace("р.", "").replace("р", "")
    t = t.replace(",", ".").strip(".")
    try:
        return abs(float(t)) or None
    except ValueError:
        return "bad"


def date(v, year):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    t = str(v).strip().split(" ")[0]
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", t)
    try:
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = re.match(r"^(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?$", t)
        if m:
            y = int(m[3]) if m[3] else year
            return dt.date(y + 2000 if y < 100 else y, int(m[2]), int(m[1]))
    except ValueError:
        pass
    return "bad"


def money(x):
    return ("%.2f" % x).rstrip("0").rstrip(".").replace(".", ",")


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["Записи"] if "Записи" in wb.sheetnames else wb.worksheets[0]
    year = dt.date.today().year
    rows, problems, last = [], [], None
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
        d, out, inc, note = (list(r) + [None] * 4)[:4]
        note = "" if note is None else str(note).strip()
        if isinstance(note, str) and note.lower().startswith("итог"):
            continue
        a_out, a_in = amount(out), amount(inc)
        if a_out is None and a_in is None:
            continue
        if "bad" in (a_out, a_in):
            problems.append(f"строка {i}: сумма не распознана ({out!r} / {inc!r})")
            continue
        day = date(d, year) if d not in (None, "") else last
        if day in (None, "bad"):
            problems.append(f"строка {i}: " + ("нет даты" if day is None else f"дата не распознана ({d!r})"))
            continue
        last = day
        rows.append((day.strftime("%d.%m.%Y"), money(a_out) if a_out else "", money(a_in) if a_in else "", note))
    with open(dst, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Дата", "Расход", "Доход", "Примечание"])
        w.writerows(rows)
    spent = sum(float(r[1].replace(",", ".")) for r in rows if r[1])
    got = sum(float(r[2].replace(",", ".")) for r in rows if r[2])
    print(f"Записей: {len(rows)} · расходы {spent:,.2f} · доходы {got:,.2f}".replace(",", " "))
    if rows:
        days = sorted(dt.datetime.strptime(r[0], "%d.%m.%Y") for r in rows)
        print(f"Даты: {days[0]:%d.%m.%Y} … {days[-1]:%d.%m.%Y}")
    for p in problems:
        print("⚠", p)
    return rows, problems


if __name__ == "__main__":
    src = sys.argv[1]
    main(src, sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.xlsx$", "", src) + "_для_DASEIN.csv")
